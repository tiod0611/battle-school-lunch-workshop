#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Excel 파일을 OpenAPI 3.0.4 JSON으로 변환
"""

import json
from pathlib import Path
from typing import Any, Dict, List

def parse_excel_files() -> List[Dict[str, Any]]:
    """Excel 파일을 파싱하고 API 정보를 추출"""
    try:
        import openpyxl
    except ImportError:
        print("openpyxl이 설치되지 않았습니다.")
        print("설치 명령: pip install -r requirements.txt")
        return []
    
    data_dir = Path(__file__).parent.parent / "data"
    excel_files = list(data_dir.glob("*_오픈API명세서.xlsx"))
    
    apis = []
    
    for excel_file in excel_files:
        print(f"파일 읽기: {excel_file.name}")
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        
        api_info = {
            "name": "",
            "path": "",
            "basic_params": [],
            "request_params": []
        }
        
        current_section = None
        params_list = []
        
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            cell_value = row[0] if row else None
            
            if not cell_value:
                continue
            
            cell_value = str(cell_value).strip()
            
            # API 이름 추출
            if "학교기본정보" in cell_value or "급식식단정보" in cell_value:
                api_info["name"] = cell_value
            
            # 요청주소 추출
            elif cell_value == "• 요청주소":
                next_cell = row[0] if row else None
                for next_row in ws.iter_rows(min_row=row_idx+1, max_row=row_idx+1, values_only=True):
                    if next_row and next_row[0]:
                        url = str(next_row[0]).strip()
                        if url.startswith("-"):
                            api_info["path"] = url[2:].strip()
            
            # 기본인자 섹션
            elif cell_value == "• 기본인자":
                current_section = "basic"
                params_list = []
            
            # 요청인자 섹션
            elif cell_value == "• 요청인자":
                if params_list:
                    if current_section == "basic":
                        api_info["basic_params"] = params_list
                    elif current_section == "request":
                        api_info["request_params"] = params_list
                current_section = "request"
                params_list = []
            
            # 파라미터 행 (변수명 | 타입 | 설명 | 설명)
            elif current_section and row[0] and row[1] and "|" not in str(row[0]):
                if not str(row[0]).startswith("•"):
                    param = {
                        "name": str(row[0]).strip(),
                        "type": str(row[1]).strip() if row[1] else "",
                        "description": str(row[2]).strip() if len(row) > 2 and row[2] else ""
                    }
                    if param["name"] not in ["변수명"]:  # 헤더 제외
                        params_list.append(param)
        
        # 마지막 섹션 저장
        if params_list and current_section == "request":
            api_info["request_params"] = params_list
        elif params_list and current_section == "basic":
            api_info["basic_params"] = params_list
        
        if api_info["path"]:
            apis.append(api_info)
        
        wb.close()
    
    return apis

def build_openapi_spec(apis: List[Dict[str, Any]]) -> Dict[str, Any]:
    """OpenAPI 3.0.4 스펙 생성"""
    spec = {
        "openapi": "3.0.4",
        "info": {
            "title": "NEIS Open API",
            "description": "학교 기본정보 및 급식 식단정보 공개 API",
            "version": "1.0.0"
        },
        "servers": [
            {
                "url": "https://open.neis.go.kr",
                "description": "NEIS Open API Server"
            }
        ],
        "paths": {},
        "components": {
            "schemas": {
                "Error": {
                    "type": "object",
                    "properties": {
                        "code": {"type": "string"},
                        "message": {"type": "string"}
                    }
                }
            }
        }
    }
    
    for api in apis:
        path_key = api["path"]
        
        params = []
        
        # 기본 파라미터
        for param in api["basic_params"]:
            params.append({
                "name": param["name"],
                "in": "query",
                "required": "필수" in param.get("type", ""),
                "description": param.get("description", ""),
                "schema": {
                    "type": get_openapi_type(param.get("type", ""))
                }
            })
        
        # 요청 파라미터
        for param in api["request_params"]:
            params.append({
                "name": param["name"],
                "in": "query",
                "required": "필수" in param.get("type", ""),
                "description": param.get("description", ""),
                "schema": {
                    "type": get_openapi_type(param.get("type", ""))
                }
            })
        
        spec["paths"][path_key] = {
            "get": {
                "summary": api["name"],
                "operationId": "getSchoolInfo" if "학교" in api["name"] else "getMealServiceDietInfo",
                "description": "공개 API 조회",
                "parameters": params,
                "responses": {
                    "200": {
                        "description": "성공적인 응답",
                        "content": {
                            "application/json": {
                                "schema": {"type": "object"}
                            },
                            "application/xml": {
                                "schema": {"type": "object"}
                            }
                        }
                    },
                    "400": {
                        "description": "요청 오류",
                        "content": {
                            "application/json": {
                                "schema": {"$ref": "#/components/schemas/Error"}
                            }
                        }
                    },
                    "500": {
                        "description": "서버 오류",
                        "content": {
                            "application/json": {
                                "schema": {"$ref": "#/components/schemas/Error"}
                            }
                        }
                    }
                }
            }
        }
    
    return spec

def get_openapi_type(excel_type: str) -> str:
    """Excel 타입을 OpenAPI 타입으로 변환"""
    if "STRING" in excel_type.upper():
        return "string"
    elif "INTEGER" in excel_type.upper():
        return "integer"
    elif "BOOLEAN" in excel_type.upper():
        return "boolean"
    return "string"

def main():
    """메인 함수"""
    print("Excel 파일 파싱 중...")
    apis = parse_excel_files()
    
    if not apis:
        print("API 정보를 추출할 수 없습니다.")
        return
    
    print(f"추출된 API 개수: {len(apis)}")
    for api in apis:
        print(f"  - {api['name']} ({api['path']})")
    
    print("\nOpenAPI 스펙 생성 중...")
    spec = build_openapi_spec(apis)
    
    output_file = Path(__file__).parent.parent / "data" / "openapi.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(spec, f, indent=2, ensure_ascii=False)
    
    print(f"✓ OpenAPI JSON 생성 완료: {output_file}")
    print(f"파일 크기: {output_file.stat().st_size} bytes")

if __name__ == "__main__":
    main()
